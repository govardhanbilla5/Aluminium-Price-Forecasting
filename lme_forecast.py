
"""
LME forecast service (backend).
 
Reads the LME price actuals from the sheet "lme_input" of the aluminium workbook, fixes outliers
(IQR + Winsorization, no cascading), fits a log-SARIMA(1,1,0)(0,1,1,12) on the last 72 months and
writes the predicted values of the next 12 months into the sheet "LME" of the SAME workbook - the
sheet the forecast engine reads.
 
Used by app/services/batch_forecast.py: refresh_lme_forecast() is called once at the start of
every batch run, so the batch always forecasts with the latest predicted LME prices.
 
Safe by design:
  * no "lme_input" sheet                      -> does nothing (feature off, old behaviour)
  * lme_input unchanged since the last run    -> does nothing (SARIMA is not re-fitted)
  * predicted values identical to the sheet   -> workbook is not re-saved
  * any error (file open in Excel, statsmodels missing, gap in months ...) is only logged;
    the batch continues with the LME sheet as it is.
"""
 
import logging
import os
import re
import shutil
import threading
import warnings
from copy import copy
from datetime import datetime
from typing import Optional
 
import numpy as np
import pandas as pd
from openpyxl import load_workbook
 
from app.data import excel_store
from app.data.excel_store import EXCEL_FILE_PATH
 
logger = logging.getLogger(__name__)
 
INPUT_SHEET = "lme_input"        # LME price actuals (model input)
LME_SHEET = "LME"                # sheet that receives the predicted values
FORECAST_MONTHS = 12             # months predicted after the last month of lme_input
LME_DECIMALS = 4                 # decimals written to the LME sheet (sheet shows $1.3822)
# >>> NEW (unit division 2204.6226)
LME_UNIT_DIVISOR = 2204.6226     # lme_input is in $/MT; predicted values are divided by this ($/lb) before
                                 # they are written to the LME sheet. Set to 1 if lme_input is already $/lb
# <<< end NEW
DATA_WINDOW_MONTHS = 72          # only the last N months of lme_input are used
# >>> NEW (accuracy)
BACKTEST_MONTHS = 12             # accuracy check: the last N months are hidden, predicted and compared
WRITE_ACCURACY = True            # True: model accuracy (100 - MAPE, %) is written into the LME sheet
ACCURACY_HEADER = "Accuracy (%)"  # header of the accuracy column in the LME sheet (created if missing)
ACCURACY_DECIMALS = 2            # decimals of the accuracy value
# <<< end NEW
MIN_MONTHS = 24                  # minimum history needed for the seasonal model
SARIMA_ORDER = (1, 1, 0)
SARIMA_SEASONAL_ORDER = (0, 1, 1, 12)
BACKUP_BEFORE_UPDATE = False      # copy of the workbook before the LME sheet is changed
MAX_BACKUPS = 0                  # only the newest N backup copies are kept (older ones are deleted)
SYNC_ACTUALS_TO_LME_SHEET = False  # True: months that are actual in lme_input also replace the old predicted
                                   # value in the LME sheet (existing LME rows only, no new rows)
DATE_COL_CANDIDATES = ("date", "month", "month (yyyy-mm)")
VALUE_COL_CANDIDATES = ("value", "lme", "lme price", "lme price ($/lb)")
 
_lock = threading.Lock()
_last_input_signature = None     # signature of the lme_input data used for the last refresh
 
 
# ---------------------------------------------------------------------------
# small helpers
# ---------------------------------------------------------------------------
 
# [OLD function] _find_sheet()
#   Finds a sheet by name, ignoring upper/lower case (lme_input / LME_Input ...). Not changed.
def _find_sheet(names, wanted: str) -> Optional[str]:
    """Case-insensitive sheet lookup (lme_input / LME_Input ...)."""
    return next((n for n in names if str(n).strip().lower() == wanted.strip().lower()), None)
 
 
# [OLD function] _find_col()
#   Finds the date / value column of lme_input by header name, else by position. Not changed.
def _find_col(df: pd.DataFrame, candidates, fallback_index: int):
    lower = {str(c).strip().lower(): c for c in df.columns}
    for cand in candidates:
        if cand in lower:
            return lower[cand]
    return df.columns[fallback_index]
 
 
# =====================================================================================
# NEW BLOCK - any date format (_MONTH_NAMES, _month_from_word, _ym, _infer_dayfirst, _parse_month).
# Used by read_lme_input() and _month_key() so lme_input / LME-sheet dates can be written in any style.
# =====================================================================================
# NEW: month name table (January..December) used to understand 'Oct', 'Sept', 'October' ...
_MONTH_NAMES = {
    "january": 1, "february": 2, "march": 3, "april": 4, "may": 5, "june": 6,
    "july": 7, "august": 8, "september": 9, "october": 10, "november": 11, "december": 12,
}
 
 
# [NEW function] _month_from_word()
#   'oct' / 'Sept' / 'OCTOBER' -> month number 10 / 9 / 10 (needs at least 3 letters). Helper of
#   _parse_month().
def _month_from_word(word: str) -> Optional[int]:
    """'oct' / 'Sept' / 'OCTOBER' -> 10 (at least 3 letters, prefix of the month name)."""
    word = word.lower()
    if len(word) < 3:
        return None
    return next((n for name, n in _MONTH_NAMES.items() if name.startswith(word)), None)
 
 
# [NEW function] _ym()
#   Builds the first day of a month (Timestamp) after checking year 1990-2100 and month 1-12, else
#   None. Helper of _parse_month().
def _ym(year: int, month: int) -> Optional[pd.Timestamp]:
    if 1990 <= year <= 2100 and 1 <= month <= 12:
        return pd.Timestamp(year=year, month=month, day=1)
    return None
 
 
# [NEW function] _infer_dayfirst()
#   Looks at the whole date column and decides how numeric dates like 01/10/2026 are read: day first
#   or month first (a part above 12 decides; all ambiguous -> day first).
def _infer_dayfirst(values) -> bool:
    """
    Numeric dates such as 01/10/2026 are ambiguous (1 Oct or 10 Jan?). Look at the whole column:
    a first part above 12 -> day first; a second part above 12 -> month first. If every value is
    ambiguous: all first parts '01' -> day first (01/Oct, 01/Nov ...), all second parts '01' ->
    month first; otherwise day first (India / Europe).
    """
    firsts, seconds = [], []
    for v in values:
        if v is None or hasattr(v, "strftime") or isinstance(v, (int, float)):
            continue
        m = re.fullmatch(r"\s*(\d{1,2})[-/.](\d{1,2})[-/.](\d{4})\s*", str(v))
        if m:
            firsts.append(int(m.group(1)))
            seconds.append(int(m.group(2)))
    if any(a > 12 for a in firsts):
        return True
    if any(b > 12 for b in seconds):
        return False
    if firsts and all(a == 1 for a in firsts) and any(b != 1 for b in seconds):
        return True
    if seconds and all(b == 1 for b in seconds) and any(a != 1 for a in firsts):
        return False
    return True
 
 
# [NEW function] _parse_month()
#   THE any-date-format reader. Turns one cell (Excel date, Excel serial 46296, 202610, 2026-10,
#   Oct'26, October 2026, 01-Oct-26, 10/2026, 01/10/2026 ...) into the first day of that month;
#   returns None if the cell is not a month (notes, totals, blanks).
def _parse_month(value, dayfirst: bool = True) -> Optional[pd.Timestamp]:
    """
    Any month / date cell -> Timestamp of the first day of that month (None if it is not a month).
    Handles: Excel date cells, Excel serial numbers (46296), 202610, 2026-10, 2026/10, 2026-10-01,
    2026-10-01 00:00:00, Oct'26, Oct 26, Oct-26, Oct_26, Oct26, OCT'26, Sept 2026, October 2026,
    Oct, 2026, 01-Oct-26, 1 Oct 2026, 26-Oct, 2026 Oct, 10/2026, 10-2026, 10-26, 01/10/2026,
    10/01/2026, 1.10.2026 ...
    """
    if value is None:
        return None
    if hasattr(value, "strftime"):                                    # datetime / date / Timestamp
        return None if pd.isna(value) else _ym(value.year, value.month)
    if isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        if pd.isna(value):
            return None
        number = float(value)
        if number.is_integer() and 190001 <= number <= 210012:         # 202610
            return _ym(int(number) // 100, int(number) % 100)
        if 20000 <= number <= 80000:                                   # Excel serial date
            d = pd.Timestamp("1899-12-30") + pd.Timedelta(days=number)
            return _ym(d.year, d.month)
        return None
 
    text = str(value).strip().strip("'\"")
    if not text or text.lower() in ("nan", "nat", "none"):
        return None
 
    if re.fullmatch(r"\d{5}(\.\d+)?", text) or re.fullmatch(r"\d{6}", text) or re.fullmatch(r"\d{8}", text):
        if len(text) == 8:                                             # 20261001
            return _ym(int(text[:4]), int(text[4:6]))
        return _parse_month(float(text))                               # 202610 or Excel serial text
 
    m = re.match(r"^(\d{4})\s*[-/._ ]\s*(\d{1,2})(?!\d)", text)         # 2026-10, 2026/10, 2026-10-01 ...
    if m:
        return _ym(int(m.group(1)), int(m.group(2)))
 
    m = re.search(r"(?<![A-Za-z])([A-Za-z]{3,9})\.?[\s,'\-_/]*(\d{4}|\d{2})(?!\d)", text)   # Oct'26 Oct 2026
    if m and _month_from_word(m.group(1)):
        year = int(m.group(2))
        return _ym(year if len(m.group(2)) == 4 else 2000 + year, _month_from_word(m.group(1)))
 
    m = re.search(r"(?<!\d)(\d{4})[\s,'\-_/]*([A-Za-z]{3,9})(?![A-Za-z])", text)           # 2026 Oct
    if m and _month_from_word(m.group(2)):
        return _ym(int(m.group(1)), _month_from_word(m.group(2)))
 
    m = re.search(r"(?<!\d)(\d{1,2})[\s,'\-_/.]*([A-Za-z]{3,9})(?![A-Za-z])[\s,'\-_/.]*(\d{4})?", text)
    if m and _month_from_word(m.group(2)):                                                  # 1 Oct 2026, 26-Oct
        if m.group(3):
            return _ym(int(m.group(3)), _month_from_word(m.group(2)))
        return _ym(2000 + int(m.group(1)), _month_from_word(m.group(2)))
 
    m = re.fullmatch(r"(\d{1,2})\s*[-/.\s]\s*(\d{4})", text)                              # 10/2026
    if m:
        return _ym(int(m.group(2)), int(m.group(1)))
 
    m = re.fullmatch(r"(\d{1,2})\s*[-/.]\s*(\d{2})", text)                                 # 10-26
    if m:
        return _ym(2000 + int(m.group(2)), int(m.group(1)))
 
    m = re.fullmatch(r"(\d{1,2})\s*[-/.]\s*(\d{1,2})\s*[-/.]\s*(\d{4})(?:\s.*)?", text)    # 01/10/2026
    if m:
        a, b, year = int(m.group(1)), int(m.group(2)), int(m.group(3))
        if a > 12:
            return _ym(year, b)
        if b > 12:
            return _ym(year, a)
        return _ym(year, b if dayfirst else a)
 
    m = re.fullmatch(r"(\d{1,2})\s*[-/.]\s*(\d{1,2})\s*[-/.]\s*(\d{2})", text)             # 01/10/26
    if m:
        a, b, year = int(m.group(1)), int(m.group(2)), 2000 + int(m.group(3))
        if a > 12:
            return _ym(year, b)
        if b > 12:
            return _ym(year, a)
        return _ym(year, b if dayfirst else a)
 
    try:                                                                  # last resort: pandas
        d = pd.to_datetime(text, errors="coerce", dayfirst=dayfirst)
        return None if pd.isna(d) else _ym(d.year, d.month)
    except Exception:
        return None
 
 
# [OLD function] _month_key()
#   Month cell -> 'YYYY-MM'. CHANGED: the old body understood only '2026-10' / '2026/10' / real
#   dates; the new body calls _parse_month(), so every format works.
#   Lines marked "NEW" were added; the "OLD (replaced)" comment lines show what was there before.
# OLD (replaced) |     if hasattr(value, "strftime"):
# OLD (replaced) |         return value.strftime("%Y-%m")
# OLD (replaced) |     m = re.match(r"^\s*(\d{4})[-/](\d{1,2})", str(value))
# OLD (replaced) |     if m and 1 <= int(m.group(2)) <= 12:
# OLD (replaced) |         return f"{m.group(1)}-{int(m.group(2)):02d}"
# OLD (replaced) |     return None
# >>> NEW (any date format)
def _month_key(value, dayfirst: bool = True) -> Optional[str]:
    """Month cell (any date format, see _parse_month) -> 'YYYY-MM'; None if it is not a month."""
    month = _parse_month(value, dayfirst)
    return month.strftime("%Y-%m") if month is not None else None
# <<< end NEW
 
 
# ---------------------------------------------------------------------------
# 1) read lme_input (robust to messy rows, same cleaning as LME_IQR_fixed.py)
# ---------------------------------------------------------------------------
 
# [OLD function] read_lme_input()
#   Reads lme_input: finds columns, reads dates, cleans values ($ and commas), keeps the last 72
#   months, checks for gaps. CHANGED: only the date line (any date format now). Everything else is
#   as before.
#   Lines marked "NEW" were added; the "OLD (replaced)" comment lines show what was there before.
def read_lme_input(path: str) -> Optional[pd.Series]:
    """Monthly LME actuals from the 'lme_input' sheet (index = month start). None if no such sheet."""
    xl = pd.ExcelFile(path)
    try:
        sheet = _find_sheet(xl.sheet_names, INPUT_SHEET)
        if sheet is None:
            return None
        raw = xl.parse(sheet).dropna(how="all")
    finally:
        xl.close()
 
    raw.columns = [str(c).strip() for c in raw.columns]
    date_col = _find_col(raw, DATE_COL_CANDIDATES, 0)
    value_col = _find_col(raw, VALUE_COL_CANDIDATES, 1)
    raw = raw[[date_col, value_col]].rename(columns={date_col: "Date", value_col: "Value"})
 
    # OLD (replaced) | raw["Date"] = pd.to_datetime(raw["Date"].astype(str).str.strip(), format="mixed", errors="coerce")
    # >>> NEW (any date format)
    dayfirst = _infer_dayfirst(raw["Date"])
    parsed = raw["Date"].map(lambda v: _parse_month(v, dayfirst))            # any date format -> month start
    skipped = raw.loc[parsed.isna() & raw["Date"].notna(), "Date"]
    if len(skipped):
        logger.warning("lme_input: %d row(s) ignored, date not understood: %s", len(skipped),
                       ", ".join(repr(v) for v in list(skipped)[:5]))
    raw["Date"] = pd.to_datetime(parsed)
    # <<< end NEW
    raw["Value"] = pd.to_numeric(
        raw["Value"].astype(str)
        .str.replace(",", "", regex=False)
        .str.replace(r"[^\d\.\-]", "", regex=True)           # strip $ and other symbols
        .str.strip(),
        errors="coerce",
    )
    df = raw.dropna(subset=["Date", "Value"]).copy()
    df["Date"] = df["Date"].dt.to_period("M").dt.to_timestamp()           # month start
    df = df.drop_duplicates(subset="Date", keep="last").sort_values("Date")
    df = df.tail(DATA_WINDOW_MONTHS).reset_index(drop=True)
    if df.empty:
        return None
 
    expected = pd.date_range(df["Date"].min(), df["Date"].max(), freq="MS")
    missing = expected.difference(pd.DatetimeIndex(df["Date"]))
    if len(missing) > 0:
        raise ValueError("lme_input has a gap - missing month(s): "
                         + ", ".join(m.strftime("%b %Y") for m in missing))
 
    ts = df.set_index("Date")["Value"]
    ts.index = pd.DatetimeIndex(ts.index, freq="MS")
    return ts
 
 
# ---------------------------------------------------------------------------
# 2) outlier fix (IQR + Winsorization, no cascading) - same logic as LME_IQR_fixed.py
# ---------------------------------------------------------------------------
 
# [OLD function] iqr_fix()
#   Outlier fix (IQR + Winsorization, non-cascading). Not changed.
def iqr_fix(ts_original: pd.Series) -> pd.Series:
    mom_pct = ts_original.pct_change().dropna() * 100
    q1 = mom_pct.quantile(0.25)
    q3 = mom_pct.quantile(0.75)
    iqr = q3 - q1
    iqr_lower = q1 - 1.5 * iqr
    iqr_upper = q3 + 1.5 * iqr
    outliers = mom_pct[(mom_pct < iqr_lower) | (mom_pct > iqr_upper)]     # detected on ORIGINAL data
 
    ts_fixed = ts_original.copy()
    for date in outliers.index:
        prev_value = ts_original.shift(1).loc[date]          # ALWAYS the original previous month
        pct_change = outliers.loc[date]
        capped = iqr_upper if pct_change > iqr_upper else iqr_lower
        ts_fixed.loc[date] = prev_value * (1 + capped / 100)
        logger.info("LME outlier %s: %+.2f%% capped to %+.2f%%", date.strftime("%b %Y"), pct_change, capped)
    return ts_fixed
 
 
# ---------------------------------------------------------------------------
# 3) SARIMA forecast (log model, back-transformed with exp)
# ---------------------------------------------------------------------------
 
# [OLD function] forecast_lme()
#   log-SARIMA(1,1,0)(0,1,1,12) -> predicts the next 12 months, exp back-transform. Not changed
#   (backtest_accuracy() reuses it).
def forecast_lme(ts_fixed: pd.Series, months: int = FORECAST_MONTHS) -> pd.Series:
    """log-SARIMA(1,1,0)(0,1,1,12) on the fixed series -> predicted prices of the next `months` months."""
    from statsmodels.tsa.statespace.sarimax import SARIMAX      # imported here: app starts even if missing
 
    if (ts_fixed <= 0).any():
        raise ValueError("lme_input has zero or negative prices - the log model cannot be used")
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        fit = SARIMAX(np.log(ts_fixed), order=SARIMA_ORDER,
                      seasonal_order=SARIMA_SEASONAL_ORDER).fit(disp=False)
        if not fit.mle_retvals.get("converged", True):
            logger.warning("LME SARIMA did not converge cleanly - forecast may be less reliable")
        return np.exp(fit.get_forecast(months).predicted_mean)
 
 
# [NEW function] backtest_accuracy()
#   MODEL ACCURACY. Same backtest as LME_IQR_fixed.py: hides the last 12 months of the fixed series,
#   fits the same model on the rest (via forecast_lme), predicts those 12 months, MAPE =
#   mean(|actual - predicted| / actual) x 100, returns accuracy = 100 - MAPE (None if less than 36
#   months of data).
def backtest_accuracy(ts_fixed: pd.Series, months: int = BACKTEST_MONTHS) -> Optional[float]:
    """
    Same backtest as LME_IQR_fixed.py: hide the last `months` months of the fixed series, fit the same
    log-SARIMA on the rest, predict those months, MAPE = mean(|actual - predicted| / actual) * 100.
    Returns accuracy = 100 - MAPE (None if there is too little data).
    """
    if len(ts_fixed) < months + MIN_MONTHS:
        logger.warning("LME accuracy skipped: %d months of data, %d needed for a %d-month backtest",
                       len(ts_fixed), months + MIN_MONTHS, months)
        return None
    train, test = ts_fixed.iloc[:-months], ts_fixed.iloc[-months:]
    predicted = forecast_lme(train, months)
    mape = float((abs(test.values - predicted.values) / test.values).mean() * 100)
    return round(100 - mape, ACCURACY_DECIMALS)
 
 
# [OLD function] _prune_backups()
#   Keeps only the newest MAX_BACKUPS backup copies of the workbook. Not changed.
def _prune_backups(path: str) -> None:
    """Keep only the newest MAX_BACKUPS files named <workbook>_backup_YYYYMMDD_HHMMSS.<ext>."""
    folder = os.path.dirname(os.path.abspath(path))
    root, ext = os.path.splitext(os.path.basename(path))
    pattern = re.compile(r"^%s_backup_\d{8}_\d{6}%s$" % (re.escape(root), re.escape(ext)))
    backups = sorted(f for f in os.listdir(folder) if pattern.match(f))
    for name in backups[:-MAX_BACKUPS]:
        try:
            os.remove(os.path.join(folder, name))
        except OSError:
            pass
 
 
# ---------------------------------------------------------------------------
# 4) write the predicted values into the LME sheet
# ---------------------------------------------------------------------------
 
# [OLD function] update_lme_sheet()
#   Writes the predicted values into the LME sheet (overwrite existing months, append new months,
#   backup, save only if something changed). CHANGED: new parameter 'accuracy', accuracy column
#   block, and dates of column A read with _month_key(v, sheet_dayfirst).
#   Lines marked "NEW" were added; the "OLD (replaced)" comment lines show what was there before.
def update_lme_sheet(path: str, forecast: pd.Series, decimals: int = LME_DECIMALS,
                     # OLD (replaced) | actuals: Optional[pd.Series] = None) -> list:
                     # >>> NEW (accuracy)
                     actuals: Optional[pd.Series] = None, accuracy: Optional[float] = None) -> list:
                     # <<< end NEW
    """
    month already in column A of the LME sheet -> value in column B is overwritten (format kept)
    month not in the sheet                     -> new row at the bottom (same formatting)
    Earlier months (actuals) and the other columns (e.g. Legend) are not touched.
    Returns the list of (month, old, new, action); the workbook is saved only if something changed.
    """
    wb = load_workbook(path, keep_vba=path.lower().endswith(".xlsm"))
    name = _find_sheet(wb.sheetnames, LME_SHEET)
    if name is None:
        raise ValueError(f"Sheet '{LME_SHEET}' not found in {path}. Sheets: {wb.sheetnames}")
    ws = wb[name]
 
    # >>> NEW (any date format)
    sheet_dayfirst = _infer_dayfirst([ws.cell(row=r, column=1).value for r in range(2, ws.max_row + 1)])
    # <<< end NEW
    month_row, last_row = {}, 1
    for r in range(2, ws.max_row + 1):
        v = ws.cell(row=r, column=1).value
        if v is None or str(v).strip() == "":
            continue
        last_row = r
        # OLD (replaced) | key = _month_key(v)
        # >>> NEW (any date format)
        key = _month_key(v, sheet_dayfirst)
        # <<< end NEW
        if key:
            month_row[key] = r
    template_row = last_row                       # formatting of the last data row is copied to new rows
    text_months = isinstance(ws.cell(row=template_row, column=1).value, str)
 
    log = []
    if actuals is not None:                        # optional: actual months replace stale predicted values
        for date, value in actuals.items():
            key = date.strftime("%Y-%m")
            if key not in month_row:
                continue
            r, new_val = month_row[key], round(float(value), decimals)
            old = ws.cell(row=r, column=2).value
            if old is None or abs(float(old) - new_val) >= 10 ** (-decimals) / 2:
                ws.cell(row=r, column=2).value = new_val
                log.append((key, old, new_val, "actual"))
 
    for date, value in forecast.items():
        key = date.strftime("%Y-%m")
        new_val = round(float(value), decimals)
        if key in month_row:
            r = month_row[key]
            old = ws.cell(row=r, column=2).value
            try:
                if old is not None and abs(float(old) - new_val) < 10 ** (-decimals) / 2:
                    continue                       # already holds this predicted value
            except (TypeError, ValueError):
                pass                               # old cell is not a number -> overwrite it
            ws.cell(row=r, column=2).value = new_val
            log.append((key, old, new_val, "updated"))
        else:
            last_row += 1
            r = last_row
            for c in (1, 2):
                src, dst = ws.cell(row=template_row, column=c), ws.cell(row=r, column=c)
                dst.font, dst.fill, dst.border = copy(src.font), copy(src.fill), copy(src.border)
                dst.alignment, dst.number_format = copy(src.alignment), src.number_format
            ws.cell(row=r, column=1).value = key if text_months else date.to_pydatetime()
            ws.cell(row=r, column=2).value = new_val
            month_row[key] = r
            log.append((key, None, new_val, "added"))
 
    # >>> NEW (accuracy)
    if accuracy is not None:                       # model accuracy (100 - MAPE) in its own column, row 2
        acc_col = next((c for c in range(1, ws.max_column + 1)
                        if str(ws.cell(row=1, column=c).value or "").strip().lower().startswith("accuracy")), None)
        if acc_col is None:                        # no such column yet -> create it (one empty column gap)
            acc_col = ws.max_column + 2
            head_src, head_dst = ws.cell(row=1, column=2), ws.cell(row=1, column=acc_col)
            head_dst.value = ACCURACY_HEADER
            head_dst.font, head_dst.fill, head_dst.border = copy(head_src.font), copy(head_src.fill), copy(head_src.border)
            head_dst.alignment = copy(head_src.alignment)
            ws.column_dimensions[head_dst.column_letter].width = max(len(ACCURACY_HEADER) + 4, 14)
        cell = ws.cell(row=2, column=acc_col)
        old_acc = cell.value
        try:
            same = old_acc is not None and abs(float(old_acc) - accuracy) < 10 ** (-ACCURACY_DECIMALS) / 2
        except (TypeError, ValueError):
            same = False
        if not same:
            cell.value = accuracy
            cell.number_format = "0." + "0" * ACCURACY_DECIMALS if ACCURACY_DECIMALS else "0"
            log.append(("Accuracy", old_acc, accuracy, "accuracy"))
    # <<< end NEW
 
    if not log:
        return log                                 # nothing to change -> no backup, no save
 
    formula_sheets = [w.title for w in wb.worksheets
                      if any(c.data_type == "f" for row in w.iter_rows() for c in row)]
    if formula_sheets:
        logger.warning("Formulas found in sheets %s: open %s in Excel once and press Save so their "
                       "values are recalculated (openpyxl does not store formula results)",
                       formula_sheets, path)
 
    backup_path = None
    if BACKUP_BEFORE_UPDATE:
        root, ext = os.path.splitext(path)
        backup_path = f"{root}_backup_{datetime.now().strftime('%Y%m%d_%H%M%S')}{ext}"
        shutil.copy2(path, backup_path)
    try:
        wb.save(path)                              # PermissionError here = file open in Excel
    except Exception:
        if backup_path and os.path.exists(backup_path):
            os.remove(backup_path)                 # nothing was changed -> do not leave a useless backup
        raise
    if BACKUP_BEFORE_UPDATE:
        _prune_backups(path)
    return log
 
 
# ---------------------------------------------------------------------------
# 5) entry point used by the batch run
# ---------------------------------------------------------------------------
 
# [OLD function] _clear_excel_store_caches()
#   Makes the engine re-read the workbook after the LME sheet was changed. Not changed.
def _clear_excel_store_caches() -> None:
    """Make the engine re-read the workbook (excel_store keeps cached copies of the sheets)."""
    try:
        excel_store._invalidate_cache()
    except Exception:
        pass
    for obj in vars(excel_store).values():         # also clear any lru_cache in excel_store
        clear = getattr(obj, "cache_clear", None)
        if callable(clear):
            clear()
 
 
# [OLD function] refresh_lme_forecast()
#   Entry point called by the batch: read lme_input -> IQR fix -> log-SARIMA -> LME sheet. CHANGED:
#   predicted values are divided by 2204.6226 ($/MT -> $/lb) and the accuracy is calculated and
#   passed to update_lme_sheet().
#   Lines marked "NEW" were added; the "OLD (replaced)" comment lines show what was there before.
def refresh_lme_forecast(path: str = EXCEL_FILE_PATH, force: bool = False) -> Optional[list]:
    """
    lme_input -> IQR fix -> log-SARIMA -> LME sheet. Never raises: problems are logged and the
    batch carries on with the LME sheet as it is. Returns the change log (None if skipped).
    """
    global _last_input_signature
    with _lock:
        try:
            ts = read_lme_input(path)
            if ts is None:
                logger.debug("No '%s' sheet in %s - LME refresh skipped", INPUT_SHEET, path)
                return None
            if len(ts) < MIN_MONTHS:
                logger.warning("lme_input has only %d months (need at least %d) - LME refresh skipped",
                               len(ts), MIN_MONTHS)
                return None
 
            signature = (len(ts), ts.index.max(), round(float(ts.sum()), 6), round(float(ts.iloc[-1]), 6))
            if not force and signature == _last_input_signature:
                return None                        # lme_input unchanged since the last refresh
 
            # OLD (replaced) | forecast = forecast_lme(iqr_fix(ts))
            # >>> NEW (accuracy)
            fixed = iqr_fix(ts)
            forecast = forecast_lme(fixed) / LME_UNIT_DIVISOR                # $/MT -> $/lb
            accuracy = None
            if WRITE_ACCURACY:
                try:
                    accuracy = backtest_accuracy(fixed)                      # 100 - MAPE of the 12-month backtest
                except Exception as exc:                                     # accuracy must never block the forecast
                    logger.warning("LME accuracy (backtest) failed - forecast is still updated: %s", exc)
            # <<< end NEW
            log = update_lme_sheet(path, forecast,
                                   # OLD (replaced) | actuals=ts if SYNC_ACTUALS_TO_LME_SHEET else None)
                                   # >>> NEW (accuracy)
                                   actuals=(ts / LME_UNIT_DIVISOR) if SYNC_ACTUALS_TO_LME_SHEET else None,
                                   accuracy=accuracy)
                                   # <<< end NEW
            _last_input_signature = signature
 
            # >>> NEW (accuracy)
            month_log = [e for e in log if e[3] != "accuracy"]
            # <<< end NEW
            if log:
                _clear_excel_store_caches()        # engine must re-read the updated LME sheet
            # >>> NEW (accuracy)
            if month_log:
            # <<< end NEW
                logger.info("LME sheet updated from lme_input: %d month(s) %s .. %s",
                            # OLD (replaced) | len(log), log[0][0], log[-1][0])
                            # >>> NEW (accuracy)
                            len(month_log), month_log[0][0], month_log[-1][0])
                            # <<< end NEW
            else:
                logger.info("LME sheet already holds the predicted values - no change")
            # >>> NEW (accuracy)
            if accuracy is not None:
                logger.info("LME model accuracy (%d-month backtest): %.*f%%", BACKTEST_MONTHS,
                            ACCURACY_DECIMALS, accuracy)
            # <<< end NEW
            return log
        except Exception as exc:                   # never break the batch because of the LME refresh
            logger.warning("LME refresh failed - batch continues with the existing LME sheet: %s", exc)
            return None
